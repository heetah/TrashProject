import io
import json
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from UI.backend.config import build_config
from UI.backend.database import Database
from UI.backend.runner import PipelineWorker
from UI.backend.service import ReviewService


def make_config(tmp_path: Path):
    allowed = tmp_path / "allowed"
    allowed.mkdir()
    config = build_config(
        env_path=tmp_path / "missing.env",
        environ={
            "UI_DATABASE_PATH": str(tmp_path / "ui.sqlite3"),
            "UI_UPLOAD_ROOT": str(tmp_path / "uploads"),
            "UI_OUTPUT_ROOT": str(tmp_path / "outputs"),
            "UI_DISCOVERY_ROOTS": str(tmp_path / "outputs"),
            "UI_ALLOWED_INPUT_ROOTS": str(allowed),
            "UI_FRONTEND_DIST": str(tmp_path / "dist"),
            "UI_LOG_ROOT": str(tmp_path / "logs"),
            "UI_EXPORT_ROOT": str(tmp_path / "exports"),
            "UI_IMPORT_EXISTING": "0",
            "UI_WORKER_ENABLED": "0",
        },
    )
    return config, allowed


def analysis_payload(events=None):
    events = [] if events is None else events
    litter_events = [event for event in events if event.get("type") == "litter"]
    return {
        "schema_version": "2.0.0",
        "video": {"file": "clip_annotated.mp4", "duration_sec": 12.5},
        "summary": {
            "litter_event_count": len(litter_events),
            "urinate_event_count": 0,
            "passed_vehicle_count": 2,
            "average_litter_confidence": 0.82 if litter_events else None,
            "detection_accuracy": None,
            "accuracy_status": "not_evaluated",
            "littering_plates": [],
            "review_required": bool(events),
        },
        "events": events,
    }


def write_existing_result(root: Path, events=None):
    root.mkdir(parents=True, exist_ok=True)
    (root / "clip_annotated.mp4").write_bytes(b"video")
    path = root / "clip_annotated_analysis.json"
    path.write_text(json.dumps(analysis_payload(events)), encoding="utf-8")
    return path


def test_config_reads_dotenv_without_overriding_exported_values(tmp_path):
    env_file = tmp_path / ".env"
    env_file.write_text("UI_PORT=6200\nUI_POLL_SECONDS=7\n", encoding="utf-8")
    config = build_config(env_path=env_file, environ={"UI_PORT": "6300"})
    assert config.port == 6300
    assert config.poll_seconds == 7


def test_folder_enqueue_is_allowlisted_and_skips_annotated_outputs(tmp_path):
    config, allowed = make_config(tmp_path)
    database = Database(config.database_path)
    database.initialize()
    service = ReviewService(config, database)
    (allowed / "a.mp4").write_bytes(b"a")
    (allowed / "b.MOV").write_bytes(b"b")
    (allowed / "a_annotated.mp4").write_bytes(b"output")

    jobs = service.enqueue_folder(allowed)

    assert [job["original_name"] for job in jobs] == ["a.mp4", "b.MOV"]
    assert all(job["status"] == "queued" for job in jobs)
    with pytest.raises(ValueError, match="允許範圍"):
        service.enqueue_folder(tmp_path)


def test_discovery_and_event_reviews_persist_across_service_instances(tmp_path):
    config, _allowed = make_config(tmp_path)
    event = {
        "type": "litter",
        "id": 3,
        "start_sec": 2.0,
        "end_sec": 3.0,
        "confidence": 0.82,
        "vehicle": None,
        "plate": None,
        "plate_confidence": None,
        "plate_status": "not_applicable",
        "attribution_status": "dustbin",
        "review_required": True,
    }
    write_existing_result(config.output_root, [event])
    database = Database(config.database_path)
    database.initialize()
    service = ReviewService(config, database)

    assert service.discover_existing() == 1
    case = service.list_cases("unreviewed")["items"][0]
    detail = service.get_case(case["id"])
    assert detail["review_units"][0]["event_key"] == "litter:3"
    service.save_review(
        case["id"], "litter:3", verdict="rejected", note="紙袋原本就在地上", reviewer="測試員"
    )

    restarted = ReviewService(config, Database(config.database_path))
    reviewed = restarted.list_cases("reviewed")["items"]
    assert len(reviewed) == 1
    saved = restarted.get_case(case["id"])["review_units"][0]["review"]
    assert saved["verdict"] == "rejected"
    assert saved["note"] == "紙袋原本就在地上"


def test_plate_correction_persists_without_overwriting_ai_or_reviewing(tmp_path):
    config, _allowed = make_config(tmp_path)
    event = {
        "type": "litter",
        "id": 7,
        "start_sec": 2.0,
        "end_sec": 3.0,
        "confidence": 0.82,
        "vehicle": "vehicle:4",
        "plate": "AI-0000",
        "plate_confidence": 0.7,
        "plate_status": "recognized",
        "attribution_status": "resolved",
        "review_required": True,
    }
    write_existing_result(config.output_root, [event])
    database = Database(config.database_path)
    database.initialize()
    service = ReviewService(config, database)
    service.discover_existing()
    job_id = service.list_cases("unreviewed")["items"][0]["id"]

    correction = service.save_plate_correction(
        job_id, "litter:7", corrected_plate="abc-1234"
    )

    assert correction["corrected_plate"] == "ABC-1234"
    restarted = ReviewService(config, Database(config.database_path))
    detail = restarted.get_case(job_id)
    unit = detail["review_units"][0]
    assert unit["event"]["plate"] == "AI-0000"
    assert unit["plate_correction"]["corrected_plate"] == "ABC-1234"
    assert unit["review"] is None
    assert detail["job"]["review_status"] == "unreviewed"

    assert restarted.delete_plate_correction(job_id, "litter:7") is True
    assert restarted.get_case(job_id)["review_units"][0]["plate_correction"] is None


def test_urinate_plate_can_be_corrected(tmp_path):
    config, _allowed = make_config(tmp_path)
    event = {
        "type": "urinate",
        "track_id": 3,
        "time_sec": 8.0,
        "start_sec": 3.0,
        "end_sec": 8.0,
        "confidence": 0.91,
        "vehicle": "scooter:8",
        "plate": None,
        "plate_status": "attempted_no_result",
    }
    write_existing_result(config.output_root, [event])
    database = Database(config.database_path)
    database.initialize()
    service = ReviewService(config, database)
    service.discover_existing()
    job_id = service.list_cases("unreviewed")["items"][0]["id"]

    correction = service.save_plate_correction(
        job_id, "urinate:3:8.0", corrected_plate="abc-5678"
    )

    assert correction["corrected_plate"] == "ABC-5678"


def test_uncertain_is_rejected_and_legacy_uncertain_remains_unreviewed(tmp_path):
    config, _allowed = make_config(tmp_path)
    event = {"type": "litter", "id": 3, "start_sec": 2.0, "end_sec": 3.0}
    write_existing_result(config.output_root, [event])
    database = Database(config.database_path)
    database.initialize()
    service = ReviewService(config, database)
    service.discover_existing()
    job_id = service.list_cases("unreviewed")["items"][0]["id"]

    with pytest.raises(ValueError, match="accepted 或 rejected"):
        service.save_review(job_id, "litter:3", verdict="uncertain")

    database.save_review(job_id, "litter:3", "uncertain", "舊資料", "舊審核者")
    detail = service.get_case(job_id)
    assert detail["job"]["review_status"] == "unreviewed"
    assert detail["job"]["reviewed_units"] == 0


def test_reviewed_export_contains_accepted_clips_and_excel(tmp_path, monkeypatch):
    config, _allowed = make_config(tmp_path)
    event = {
        "type": "urinate",
        "track_id": 3,
        "time_sec": 8.0,
        "start_sec": 3.0,
        "end_sec": 8.0,
        "confidence": 0.91,
        "vehicle": "scooter:8",
        "plate": "AI-0000",
        "plate_confidence": 0.82,
        "plate_status": "recognized",
        "attribution_status": "resolved",
    }
    write_existing_result(config.output_root, [event])
    database = Database(config.database_path)
    database.initialize()
    service = ReviewService(config, database)
    service.discover_existing()
    job_id = service.list_cases("unreviewed")["items"][0]["id"]
    event_key = "urinate:3:8.0"
    service.save_plate_correction(job_id, event_key, corrected_plate="HUM-1234")
    service.save_review(
        job_id,
        event_key,
        verdict="accepted",
        note="=HYPERLINK(\"bad\")",
        reviewer="測試員",
    )

    commands = []

    def fake_ffmpeg(command, *, stdout, stderr, text, check):
        commands.append(command)
        Path(command[-1]).write_bytes(b"clip")
        return SimpleNamespace(returncode=0, stderr="")

    monkeypatch.setattr("UI.backend.exporter.subprocess.run", fake_ffmpeg)
    archive_path = service.export_reviewed_violations()

    assert archive_path.parent == config.export_root
    with zipfile.ZipFile(archive_path) as archive:
        names = archive.namelist()
        assert "違規清單.xlsx" in names
        clips = [name for name in names if name.startswith("clips/")]
        assert len(clips) == 1
        workbook = load_workbook(io.BytesIO(archive.read("違規清單.xlsx")))
    sheet = workbook["違規清單"]
    headers = [cell.value for cell in sheet[1]]
    assert sheet.cell(2, headers.index("違規類型") + 1).value == "隨地便溺"
    assert sheet.cell(2, headers.index("車牌") + 1).value == "HUM-1234"
    assert sheet.cell(2, headers.index("車牌來源") + 1).value == "人工修正"
    assert sheet.cell(2, headers.index("備註") + 1).value.startswith("'")
    assert commands[0][commands[0].index("-ss") + 1] == "0.000"


def test_no_event_video_still_requires_human_review(tmp_path):
    config, _allowed = make_config(tmp_path)
    write_existing_result(config.output_root, [])
    database = Database(config.database_path)
    database.initialize()
    service = ReviewService(config, database)
    service.discover_existing()

    case = service.list_cases("unreviewed")["items"][0]
    unit = service.get_case(case["id"])["review_units"][0]
    assert unit["kind"] == "no_ai_event"
    service.save_review(case["id"], unit["event_key"], verdict="accepted")
    assert service.list_cases("reviewed")["counts"]["reviewed"] == 1


class FakeUpload:
    filename = "巷口 01.mp4"

    def save(self, destination):
        Path(destination).write_bytes(io.BytesIO(b"upload").read())


def test_uploaded_video_uses_private_job_directory(tmp_path):
    config, _allowed = make_config(tmp_path)
    database = Database(config.database_path)
    database.initialize()
    service = ReviewService(config, database)

    job = service.enqueue_upload(FakeUpload())

    saved = Path(job["source_path"])
    assert saved.is_file()
    assert saved.parent.parent == config.upload_root
    assert job["original_name"] == "巷口 01.mp4"


def test_worker_calls_positional_main_and_records_expected_outputs(tmp_path, monkeypatch):
    config, allowed = make_config(tmp_path)
    source = allowed / "clip.mp4"
    source.write_bytes(b"input")
    database = Database(config.database_path)
    database.initialize()
    service = ReviewService(config, database)
    queued = service.enqueue_path(source)
    claimed = database.claim_next_job()
    worker = PipelineWorker(config, database)

    def fake_run(command, *, cwd, env, stdout, stderr, text, check):
        assert command[-2:] == [str(config.pipeline_entrypoint), str(source)]
        assert cwd == config.repository_root
        assert env["PIPELINE_BATCH"] == "8"
        output_dir = Path(env["OUTPUT_ROOT"])
        output_dir.mkdir(parents=True, exist_ok=True)
        (output_dir / "clip_annotated.mp4").write_bytes(b"video")
        (output_dir / "clip_annotated_analysis.json").write_text(
            json.dumps(analysis_payload([])), encoding="utf-8"
        )
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr("UI.backend.runner.subprocess.run", fake_run)
    worker._run_job(claimed)

    completed = database.get_job(queued["id"])
    assert completed["status"] == "completed"
    assert Path(completed["output_video"]).name == "clip_annotated.mp4"
    assert Path(completed["analysis_path"]).name == "clip_annotated_analysis.json"
