"""Create a reviewed-violation ZIP with MP4 clips and an Excel index."""

from __future__ import annotations

import math
import os
import re
import shutil
import subprocess
import tempfile
import uuid
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


EXPORT_HEADERS = (
    "序號",
    "案件 ID",
    "原始影片",
    "違規類型",
    "事件識別",
    "開始秒數",
    "結束秒數",
    "模型信心",
    "關聯車輛",
    "車牌",
    "車牌來源",
    "OCR 信心",
    "歸因狀態",
    "人工判定",
    "審核者",
    "審核時間",
    "備註",
    "片段檔名",
    "匯出狀態",
)


def _safe_filename(value: str, fallback: str) -> str:
    stem = Path(str(value or "")).stem.strip()
    stem = re.sub(r"[^\w.()\- ]+", "_", stem, flags=re.UNICODE).strip(" ._")
    return (stem[:80] or fallback)


def _excel_text(value: Any) -> Any:
    """Keep reviewer-controlled text from becoming an Excel formula."""
    if not isinstance(value, str):
        return value
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def _format_timestamp(seconds: float) -> str:
    total = max(0, int(round(seconds)))
    hours, remainder = divmod(total, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}-{minutes:02d}-{secs:02d}"


def _finite_float(value: Any) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _cut_clip(
    source: Path,
    destination: Path,
    *,
    start_sec: float,
    end_sec: float,
    ffmpeg_executable: str,
) -> None:
    duration = max(0.1, end_sec - start_sec)
    command = [
        ffmpeg_executable,
        "-hide_banner",
        "-loglevel",
        "error",
        "-y",
        "-ss",
        f"{start_sec:.3f}",
        "-i",
        str(source),
        "-t",
        f"{duration:.3f}",
        "-map",
        "0:v:0",
        "-c:v",
        "libx264",
        "-preset",
        "veryfast",
        "-crf",
        "23",
        "-an",
        "-movflags",
        "+faststart",
        str(destination),
    ]
    try:
        result = subprocess.run(
            command,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            check=False,
        )
    except OSError as exc:
        raise ValueError(f"無法執行 FFmpeg：{exc}") from exc
    if result.returncode != 0 or not destination.is_file():
        detail = (result.stderr or "未知錯誤").strip().splitlines()
        message = detail[-1] if detail else "未知錯誤"
        raise ValueError(f"違規片段剪輯失敗：{message}")


def _write_workbook(rows: list[dict[str, Any]], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "違規清單"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_HEADERS))}1"
    header_fill = PatternFill("solid", fgColor="1D6751")
    header_font = Font(color="FFFFFF", bold=True)
    for column, header in enumerate(EXPORT_HEADERS, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, item in enumerate(rows, start=2):
        values = (
            item["sequence"],
            item["job_id"],
            item["original_name"],
            item["violation_type"],
            item["event_key"],
            item.get("event_start_sec"),
            item.get("event_end_sec"),
            item.get("confidence"),
            item.get("vehicle") or "NULL",
            item.get("plate") or "未取得",
            item.get("plate_source"),
            item.get("plate_confidence"),
            item.get("attribution_status") or "未提供",
            "違規成立",
            item.get("reviewer"),
            item.get("reviewed_at"),
            item.get("note"),
            item.get("clip_name"),
            item.get("export_status"),
        )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column, value=_excel_text(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in (6, 7):
            sheet.cell(row=row_index, column=column).number_format = "0.00"
        for column in (8, 12):
            sheet.cell(row=row_index, column=column).number_format = "0.00%"

    widths = (8, 28, 28, 14, 24, 12, 12, 12, 18, 16, 14, 12, 18, 14, 16, 24, 36, 42, 28)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    workbook.save(path)


def build_reviewed_export(
    violations: list[dict[str, Any]],
    *,
    export_root: Path,
    ffmpeg_executable: str,
    pre_roll_sec: float,
    post_roll_sec: float,
) -> Path:
    """Write one ZIP containing ``clips/*.mp4`` and ``違規清單.xlsx``."""
    if not violations:
        raise ValueError("已審核案件中沒有人工判定為違規成立的事件")

    export_root.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S")
    archive_name = f"reviewed_violations_{stamp}_{uuid.uuid4().hex[:8]}.zip"
    archive_path = export_root / archive_name
    temporary_archive = export_root / f".{archive_name}.tmp"
    staging = Path(tempfile.mkdtemp(prefix="reviewed-export-", dir=export_root))
    clips_dir = staging / "clips"
    clips_dir.mkdir()

    rows: list[dict[str, Any]] = []
    try:
        for sequence, violation in enumerate(violations, start=1):
            row = dict(violation)
            row["sequence"] = sequence
            source = Path(str(violation.get("source_video") or ""))
            event_start = _finite_float(violation.get("event_start_sec"))
            event_end = _finite_float(violation.get("event_end_sec"))
            clip_name = ""
            export_status = ""

            if not source.is_file():
                export_status = "未剪輯：找不到 annotated MP4"
            elif event_start is None and event_end is None:
                export_status = "未剪輯：事件沒有時間戳"
            else:
                start_value = event_start if event_start is not None else event_end
                end_value = event_end if event_end is not None else event_start
                if end_value < start_value:
                    start_value, end_value = end_value, start_value
                duration_value = _finite_float(violation.get("video_duration_sec"))
                clip_start = max(0.0, start_value - float(pre_roll_sec))
                clip_end = end_value + float(post_roll_sec)
                if duration_value is not None and duration_value > 0:
                    clip_end = min(clip_end, duration_value)
                if clip_end <= clip_start:
                    clip_end = min(
                        clip_start + 0.1,
                        duration_value if duration_value and duration_value > clip_start else clip_start + 0.1,
                    )
                event_label = _safe_filename(violation.get("event_key", ""), "event")
                video_label = _safe_filename(violation.get("original_name", ""), "video")
                clip_name = (
                    f"{sequence:04d}_{video_label}_{event_label}_"
                    f"{_format_timestamp(clip_start)}_{_format_timestamp(clip_end)}.mp4"
                )
                _cut_clip(
                    source,
                    clips_dir / clip_name,
                    start_sec=clip_start,
                    end_sec=clip_end,
                    ffmpeg_executable=ffmpeg_executable,
                )
                export_status = "已剪輯"

            row["clip_name"] = f"clips/{clip_name}" if clip_name else ""
            row["export_status"] = export_status
            rows.append(row)

        workbook_path = staging / "違規清單.xlsx"
        _write_workbook(rows, workbook_path)
        with zipfile.ZipFile(temporary_archive, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.write(workbook_path, workbook_path.name)
            for clip in sorted(clips_dir.glob("*.mp4")):
                archive.write(clip, f"clips/{clip.name}")
        os.replace(temporary_archive, archive_path)
        return archive_path
    finally:
        shutil.rmtree(staging, ignore_errors=True)
        if temporary_archive.exists():
            temporary_archive.unlink()
